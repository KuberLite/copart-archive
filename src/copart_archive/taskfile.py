"""Task file sent by the client: which lots need photos.

The client prepares it in Excel, so the format is loose: xlsx or csv
(possibly re-saved by Russian Excel with ';' and cp1251), extra sheets,
duplicates, deleted columns. Only a lot number is required, taken from
"Lot #" or parsed from "Lot URL".
"""

import csv
import io
import re
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

import openpyxl

MAIN_SHEET = "ALL_LOTS"
LOT_COLUMN = "lot #"
URL_COLUMN = "lot url"
_LOT_RE = re.compile(r"\d{6,9}")
_URL_RE = re.compile(r"/lot/(\d{6,9})")


class TaskFileError(ValueError):
    pass


@dataclass
class TaskFile:
    path: Path
    sheet: str | None
    lots: list[str] = field(default_factory=list)  # unique, in file order
    rows: dict[str, dict[str, str]] = field(default_factory=dict)  # first row per lot
    total_rows: int = 0
    duplicates: int = 0
    invalid: list[tuple[int, str]] = field(default_factory=list)  # (line, value)

    def report(self) -> str:
        where = f", лист {self.sheet}" if self.sheet else ""
        lines = [f"Файл {self.path.name}{where}: строк {self.total_rows}, лотов {len(self.lots)}"]
        if self.duplicates:
            lines.append(f"  повторов: {self.duplicates}")
        if self.invalid:
            sample = ", ".join(f"стр. {n}: {v!r}" for n, v in self.invalid[:5])
            lines.append(f"  без номера лота: {len(self.invalid)} ({sample})")
        return "\n".join(lines)


def _cell(value) -> str:
    """Excel cell to the text Copart would have written in the CSV."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, datetime):
        return value.strftime("%m/%d/%Y %I:%M %p").lower()
    return str(value).strip()


def _read_xlsx(path: Path) -> tuple[str, list[list[str]]]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb[MAIN_SHEET] if MAIN_SHEET in wb.sheetnames else wb.worksheets[0]
        return ws.title, [[_cell(v) for v in row] for row in ws.iter_rows(values_only=True)]
    finally:
        wb.close()


def _read_csv(path: Path) -> list[list[str]]:
    data = path.read_bytes()
    for encoding in ("utf-8-sig", "cp1251"):
        try:
            text = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    header = text.splitlines()[0] if text else ""
    delimiter = max(",;\t", key=header.count)
    return [[c.strip() for c in row] for row in csv.reader(io.StringIO(text), delimiter=delimiter)]


def _lot_number(lot_value: str, url_value: str) -> str | None:
    if _LOT_RE.fullmatch(lot_value):
        return lot_value
    m = _URL_RE.search(url_value)
    return m[1] if m else None


def read(path: Path) -> TaskFile:
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        sheet, table = _read_xlsx(path)
    elif path.suffix.lower() in (".csv", ".txt"):
        sheet, table = None, _read_csv(path)
    else:
        raise TaskFileError(f"{path.name}: нужен .xlsx или .csv")

    if not table:
        raise TaskFileError(f"{path.name}: файл пустой")
    header = table[0]
    keys = [h.strip().lower() for h in header]
    if LOT_COLUMN not in keys and URL_COLUMN not in keys:
        raise TaskFileError(f"{path.name}: нет колонки «Lot #» или «Lot URL»")
    lot_i = keys.index(LOT_COLUMN) if LOT_COLUMN in keys else None
    url_i = keys.index(URL_COLUMN) if URL_COLUMN in keys else None

    task = TaskFile(path=path, sheet=sheet)
    for line, values in enumerate(table[1:], start=2):
        if not any(values):
            continue  # Excel leaves empty rows at the end
        task.total_rows += 1
        get = lambda i: values[i] if i is not None and i < len(values) else ""
        lot = _lot_number(get(lot_i), get(url_i))
        if lot is None:
            task.invalid.append((line, get(lot_i) or get(url_i)))
        elif lot in task.rows:
            task.duplicates += 1
        else:
            task.lots.append(lot)
            task.rows[lot] = dict(zip((h.strip() for h in header), values))
    return task
