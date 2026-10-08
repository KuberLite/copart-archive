"""Reading xlsx and csv files that clients prepare by hand.

The formats are loose: Excel may store numbers as numbers, re-save csv with
';' and cp1251, or export UTF-16 "Unicode Text". Everything lands here as
rows of strings.
"""

import csv
import io
from datetime import datetime
from pathlib import Path

import openpyxl

XLSX_SUFFIXES = (".xlsx", ".xlsm")
CSV_SUFFIXES = (".csv", ".txt")


class FormatError(ValueError):
    pass


def cell(value) -> str:
    """Excel cell to the text Copart would have written in a csv."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, datetime):
        return value.strftime("%m/%d/%Y %I:%M %p").lower()
    return str(value).strip()


def read_xlsx(path: Path, sheet: str | None = None) -> list[tuple[str, list[list[str]]]]:
    """Sheets as (title, rows). With `sheet`, only that one if it exists."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        sheets = [wb[sheet]] if sheet and sheet in wb.sheetnames else wb.worksheets
        return [(ws.title, [[cell(v) for v in row] for row in ws.iter_rows(values_only=True)])
                for ws in sheets]
    finally:
        wb.close()


def decode(path: Path) -> str:
    data = path.read_bytes()
    # Excel's "Unicode Text" is UTF-16 with a BOM
    encodings = ("utf-16",) if data[:2] in (b"\xff\xfe", b"\xfe\xff") else ("utf-8-sig", "cp1251")
    for encoding in encodings:
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise FormatError(f"{path.name}: не удалось определить кодировку, сохраните как CSV UTF-8")


def read_csv(path: Path) -> list[list[str]]:
    text = decode(path)
    header = text.splitlines()[0] if text else ""
    delimiter = max(",;\t", key=header.count)
    return [[c.strip() for c in row] for row in csv.reader(io.StringIO(text), delimiter=delimiter)]


def read_headers(path: Path) -> list[tuple[str | None, list[str]]]:
    """Only the first row of each table — enough to tell one file format from another
    without parsing a workbook of a couple of hundred sheets."""
    if path.suffix.lower() in XLSX_SUFFIXES:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            return [(ws.title, [cell(v) for v in next(ws.iter_rows(max_row=1, values_only=True), ())])
                    for ws in wb.worksheets]
        finally:
            wb.close()
    if path.suffix.lower() in CSV_SUFFIXES:
        rows = read_csv(path)
        return [(None, rows[0] if rows else [])]
    raise FormatError(f"{path.name}: нужен .xlsx или .csv")


def read_tables(path: Path, sheet: str | None = None) -> list[tuple[str | None, list[list[str]]]]:
    """Tables of a file: one per sheet for xlsx, a single one for csv."""
    if path.suffix.lower() in XLSX_SUFFIXES:
        return list(read_xlsx(path, sheet))
    if path.suffix.lower() in CSV_SUFFIXES:
        return [(None, read_csv(path))]
    raise FormatError(f"{path.name}: нужен .xlsx или .csv")
